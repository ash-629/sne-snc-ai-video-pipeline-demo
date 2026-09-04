#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""SNE (Simply Natural English) episode docx generator — English only.

Input: script.txt (one `Speaker: line` per turn, blank line between turns).
Output: A4 .docx with title block, speaker legend, color-coded dialogue,
        and a "Key Vocabulary" section. Everything is English.
"""
import argparse
import json
import re
import sys
from pathlib import Path

if sys.stdout.encoding and sys.stdout.encoding.lower() in ('gbk', 'cp936', 'gb2312'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


# SNE brand-adjacent palette (slightly cooler than SNC's orange)
COLOR_TITLE   = RGBColor(0x1F, 0x6F, 0xB4)   # #1F6FB4 deep blue — main accent
COLOR_GRAY    = RGBColor(0x4A, 0x4A, 0x4A)
COLOR_BLUE    = RGBColor(0x2E, 0x75, 0xB6)
COLOR_TEAL    = RGBColor(0x1A, 0x7A, 0x6E)
COLOR_ORANGE  = RGBColor(0xC0, 0x5C, 0x1A)
COLOR_BLACK   = RGBColor(0x1A, 0x1A, 0x1A)

SPEAKER_COLORS = {
    'Maya':  {'bg': RGBColor(0xE8, 0xF5, 0xF3), 'border': COLOR_TEAL},
    'Ethan': {'bg': RGBColor(0xFE, 0xF3, 0xEC), 'border': COLOR_ORANGE},
}
_FALLBACK = [
    {'bg': RGBColor(0xE8, 0xF5, 0xF3), 'border': COLOR_TEAL},
    {'bg': RGBColor(0xFE, 0xF3, 0xEC), 'border': COLOR_ORANGE},
    {'bg': RGBColor(0xEE, 0xEE, 0xF6), 'border': COLOR_TITLE},
]
_fallback_idx = {}

ENGLISH_FONT = 'Calibri'


def _speaker_color(name):
    if name in SPEAKER_COLORS:
        return SPEAKER_COLORS[name]
    if name not in _fallback_idx:
        _fallback_idx[name] = len(_fallback_idx) % len(_FALLBACK)
    return _FALLBACK[_fallback_idx[name]]


def parse_script(path):
    entries = []
    text = Path(path).read_text(encoding='utf-8')
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        m = re.match(r'^([^:：]+)[:：]\s*(.+)$', line)
        if m:
            entries.append({'speaker': m.group(1).strip(), 'text': m.group(2).strip()})
    return entries


def _set_run_font(run, size_pt, color_rgb, bold=False, italic=False):
    run.font.name = ENGLISH_FONT
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color_rgb
    run.font.bold = bold
    run.font.italic = italic


def _add_dialogue_block(doc, speaker, text, color):
    border_hex = ''.join(f'{c:02X}' for c in color['border'])
    bg_hex = ''.join(f'{c:02X}' for c in color['bg'])

    # Speaker label paragraph
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent = Cm(0.5)
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:color="{border_hex}"/></w:pBdr>'
    ))
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{bg_hex}" w:val="clear"/>'
    ))
    _set_run_font(p.add_run(speaker), 10, color['border'], bold=True)

    # Text paragraph
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent = Cm(0.5)
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:color="{border_hex}"/></w:pBdr>'
    ))
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{bg_hex}" w:val="clear"/>'
    ))
    _set_run_font(p.add_run(text), 12, COLOR_BLACK)


def _add_hr(doc, color_hex, sz=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:color="{color_hex}"/></w:pBdr>'
    ))


def generate(script_path, output_path, title, subtitle, vocab_list,
             channel_name='Simply Natural English',
             section_label='Casual English · Phrase Deep-Dive'):
    entries = parse_script(script_path)
    print(f'  parsed {len(entries)} dialogue turns', flush=True)

    doc = Document()
    for s in doc.sections:
        s.page_width = Cm(21)
        s.page_height = Cm(29.7)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(2.0)
        s.top_margin = Cm(1.8)
        s.bottom_margin = Cm(1.8)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run(title), 36, COLOR_TITLE, bold=True)

    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(p.add_run(subtitle), 14, COLOR_GRAY, italic=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run(channel_name), 11, COLOR_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run(section_label), 10, COLOR_GRAY)

    _add_hr(doc, '1F6FB4', sz=8)

    # Speaker legend
    seen = []
    for e in entries:
        if e['speaker'] not in seen:
            seen.append(e['speaker'])
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    for i, spk in enumerate(seen):
        col = _speaker_color(spk)
        if i > 0:
            p.add_run('   ')
        _set_run_font(p.add_run(f'■ {spk}'), 11, col['border'], bold=True)
    _add_hr(doc, 'F0F0F0', sz=4)

    # Dialogue
    for e in entries:
        _add_dialogue_block(doc, e['speaker'], e['text'], _speaker_color(e['speaker']))

    # Vocabulary section
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run('Key Vocabulary & Phrases'), 16, COLOR_TITLE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(
        p.add_run('The casual building blocks you actually need to use this phrase like a native.'),
        10, COLOR_GRAY, italic=True,
    )
    doc.add_paragraph()

    for i, v in enumerate(vocab_list, 1):
        # Headword line: "1. word — part of speech"
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        _set_run_font(p.add_run(f'{i}. '), 12, COLOR_GRAY, bold=True)
        _set_run_font(p.add_run(v['word']), 13, COLOR_TITLE, bold=True)
        pos = v.get('pos')
        if pos:
            _set_run_font(p.add_run(f'   ({pos})'), 10, COLOR_GRAY, italic=True)

        # Definition
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(2)
        _set_run_font(p.add_run(v['definition']), 11, COLOR_BLACK)

        # Example (italic, blue)
        if v.get('example'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(2)
            _set_run_font(p.add_run('e.g. '), 10, COLOR_GRAY, italic=True)
            _set_run_font(p.add_run(v['example']), 10, COLOR_BLUE, italic=True)

        # Optional usage note
        if v.get('note'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(8)
            _set_run_font(p.add_run('Note: '), 10, COLOR_ORANGE, bold=True)
            _set_run_font(p.add_run(v['note']), 10, COLOR_GRAY)
        else:
            doc.add_paragraph()  # gap

    doc.save(output_path)
    print(f'  wrote {output_path}', flush=True)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--script', required=True, help='script.txt path')
    ap.add_argument('--output', required=True, help='output .docx path')
    ap.add_argument('--title', required=True, help='main title (English)')
    ap.add_argument('--subtitle', default=None, help='optional italic subtitle')
    ap.add_argument('--channel-name', default='Simply Natural English')
    ap.add_argument('--section-label', default='Casual English · Phrase Deep-Dive')
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument('--vocab-file')
    g.add_argument('--vocab-json')
    args = ap.parse_args()

    if args.vocab_file:
        vocab = json.loads(Path(args.vocab_file).read_text(encoding='utf-8'))
    else:
        vocab = json.loads(args.vocab_json)
    print(f'  loaded {len(vocab)} vocab items', flush=True)

    generate(args.script, args.output, args.title, args.subtitle, vocab,
             channel_name=args.channel_name, section_label=args.section_label)
